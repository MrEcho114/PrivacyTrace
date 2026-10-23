#!/usr/bin/env python
"""Assemble the rendered academic-paper deck into a PPTX and record its provenance.

Pipeline (run from the repository root):
    node scripts/render-academic-deck.cjs      # HTML -> deck/render/slide-*.png
    python showcase/privacytrace-review/deck/build_academic_pptx.py

The PPTX keeps one 1920x1080 rendered slide per page plus the speaker notes from
the HTML source. ``delivery/academic/manifest.json`` records the source HTML,
every render hash, the deck hash/bytes and the capture batch, so
``scripts/verify-academic-deck.cjs`` can re-read the PPTX pictures and prove they
are the renders the manifest names.
"""
from __future__ import annotations

import hashlib
import html as html_lib
import json
import re
import zipfile
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

DECK = Path(__file__).resolve().parent
MATERIAL = DECK.parent
SOURCE = DECK / "academic.html"
RENDER = DECK / "render"
OUT_DIR = MATERIAL / "delivery" / "academic"
OUT = OUT_DIR / "PrivacyTrace-Review-Academic.pptx"
MANIFEST = OUT_DIR / "manifest.json"
CAPTURE_MANIFEST = MATERIAL / "capture" / "manifest.json"

REQUIRED_REFS = (
    "../capture/mobile-import-390x844-dpr3.png",
    "../capture/mobile-report-390x844-dpr3.png",
)
BANNED = ("Issue #", "未上传", "最终上传未发生", "2026-10-22", "07b37b2", "工作台 · 开发版", "交付候选")


def sha256(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def normalize_package(path: Path, slides: int, notes: int) -> None:
    """Pin the ZIP timestamps and replace python-pptx template metadata."""
    stamp = (2026, 10, 23, 0, 0, 0)
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "docProps/app.xml":
                text = data.decode("utf-8")
                text = re.sub(r"<Application>.*?</Application>", "<Application>PrivacyTrace</Application>", text)
                text = re.sub(r"<PresentationFormat>.*?</PresentationFormat>", "<PresentationFormat>On-screen Show (16:9)</PresentationFormat>", text)
                text = re.sub(r"<Slides>\d+</Slides>", f"<Slides>{slides}</Slides>", text)
                text = re.sub(r"<Notes>\d+</Notes>", f"<Notes>{notes}</Notes>", text)
                text = re.sub(r"<AppVersion>.*?</AppVersion>", "<AppVersion>16.0000</AppVersion>", text)
                data = text.encode("utf-8")
            info.date_time = stamp
            target.writestr(info, data)
    tmp.replace(path)


def main() -> int:
    html = SOURCE.read_text(encoding="utf-8")
    notes = [html_lib.unescape(re.sub(r"<[^>]+>", "", block)).strip()
             for block in re.findall(r'<div class="notes">(.*?)</div>', html, re.S)]
    if len(notes) != 10:
        raise SystemExit(f"expected 10 speaker-notes blocks, found {len(notes)}")
    if html.count('class="slide"') != 10:
        raise SystemExit(f"expected 10 slides in source, found {html.count('class=\"slide\"')}")
    for ref in REQUIRED_REFS:
        if ref not in html:
            raise SystemExit(f"source does not reference the capture screenshot {ref}")
    banned = [phrase for phrase in BANNED if phrase in html]
    if banned:
        raise SystemExit(f"banned strings in the deck source: {banned}")

    images = sorted(RENDER.glob("slide-*.png"))
    if len(images) != 10:
        raise SystemExit(f"expected 10 rendered slides, found {len(images)} (run scripts/render-academic-deck.cjs)")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for image, note in zip(images, notes):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(image), 0, 0, width=prs.slide_width, height=prs.slide_height)
        slide.notes_slide.notes_text_frame.text = note

    core = prs.core_properties
    core.title = "PrivacyTrace · 作品评审材料（学术论文风）"
    core.author = "PrivacyTrace"
    core.last_modified_by = "PrivacyTrace"
    core.subject = "10 页评审 PPT · 内容与已对齐版一致 · 视觉采用 hps-academic-paper 模板"
    core.comments = ""
    core.keywords = ""
    core.category = ""
    core.revision = 1
    core.created = datetime(2026, 10, 23, 0, 0, 0)
    core.modified = datetime(2026, 10, 23, 0, 0, 0)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    normalize_package(OUT, len(prs.slides), len(notes))

    capture = json.loads(CAPTURE_MANIFEST.read_text(encoding="utf-8"))
    screens = []
    for screen in capture["screens"]:
        if screen["file"] not in {"mobile-import-390x844-dpr3.png", "mobile-report-390x844-dpr3.png"}:
            continue
        screens.append({
            "role": "import" if "import" in screen["file"] else "report",
            "file": f"../../capture/{screen['file']}",
            "bytes": screen["bytes"],
            "sha256": screen["sha256"],
            "viewport": screen["viewport"],
            "device_scale_factor": screen["device_scale_factor"],
            "captured_at": screen["captured_at"],
        })

    manifest = {
        "schema": "privacytrace/academic-deck@1",
        "title": "PrivacyTrace · 作品评审材料（学术论文风）",
        "version": "0.1",
        "kind": "academic-paper",
        "template": {
            "id": "od-hps-academic-paper",
            "license": "MIT",
            "attribution": "lewislulu/html-ppt-skill · theme academic-paper",
        },
        "source": {
            "html": "../../deck/academic.html",
            "html_sha256": sha256(SOURCE),
            "slides": 10,
            "notes": 10,
        },
        "render": {
            "width": 1920,
            "height": 1080,
            "slides": [
                {"file": image.name, "bytes": image.stat().st_size, "sha256": sha256(image)}
                for image in images
            ],
        },
        "deck": {
            "file": OUT.name,
            "bytes": OUT.stat().st_size,
            "sha256": sha256(OUT),
            "slides": len(prs.slides),
            "notes": len(notes),
        },
        "capture": {
            "manifest": "../../capture/manifest.json",
            "manifest_sha256": sha256(CAPTURE_MANIFEST),
            "source_commit": capture["source_commit"],
            "source_commit_short": capture["source_commit_short"],
            "capture_date": capture["capture_date"],
            "screens": screens,
        },
        "captions": {
            "client": "图 3 / 图 4：新版 Android App 导入与报告",
            "capture": f"{capture['capture_date']} · 390×844 DPR 3 · {capture['source_commit_short']}",
        },
        "banned_hits": [],
    }
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    print("wrote", OUT, manifest["deck"]["bytes"], "bytes", manifest["deck"]["sha256"])
    print("wrote", MANIFEST)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
